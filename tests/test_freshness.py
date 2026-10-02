import copy

import pytest

from scripts import check_freshness as cf

FACTS = cf.load_facts()
TER = cf.parse_ter("data/raw/S10.xlsx")           # the committed copies of the files facts.json was built from
RISK = cf.parse_riskometer("data/raw/S11.xlsx")


def test_committed_files_parse_and_match_facts_json():
    assert len(TER) == 4 and len(RISK) == 4
    assert cf.compare(FACTS, TER, RISK) == []


def test_changed_ter_value_is_reported_with_both_numbers():
    ter = copy.deepcopy(TER)
    ter["Groww Large Cap Fund"]["direct"] = 1.75
    p = cf.compare(FACTS, ter, RISK)
    assert len(p) == 1 and "Groww Large Cap Fund Direct plan is 1.75%" in p[0] and "1.69%" in p[0]


def test_tiny_rounding_noise_is_not_a_change():
    ter = copy.deepcopy(TER)
    ter["Groww Multicap Fund"]["regular"] = round(ter["Groww Multicap Fund"]["regular"] + 0.004, 3)
    assert cf.compare(FACTS, ter, RISK) == []


def test_changed_riskometer_is_reported():
    risk = dict(RISK)
    risk["Groww Small Cap Fund"] = ("Sep 2026", "High")
    p = cf.compare(FACTS, TER, risk)
    assert len(p) == 1 and "Small Cap" in p[0] and "'High'" in p[0] and "Sep 2026" in p[0]


def test_empty_parse_is_a_problem_never_a_silent_pass():
    assert any("missing from the live file" in x for x in cf.compare(FACTS, {}, RISK))
    assert any("missing from the live file" in x for x in cf.compare(FACTS, TER, {}))


def test_none_means_file_not_fetched_and_is_skipped_not_duplicated():
    assert cf.compare(FACTS, None, None) == []


def test_latest_ter_file_is_chosen_by_date():
    html = ('"https://assets-netstorage.growwmf.in/compliance_docs/Downloads/Expense%20Ratio/2026/September/GRMF_WEBSITE%20TER_29.09.2026.xlsx" '
            '"https://assets-netstorage.growwmf.in/compliance_docs/Downloads/Expense%20Ratio/2026/October/GRMF_WEBSITE%20TER_01.10.2026.xlsx" '
            '"https://assets-netstorage.growwmf.in/compliance_docs/Downloads/Expense%20Ratio/2026/September/GRMF_WEBSITE%20TER_30.09.2026.xlsx"')
    url, d = cf.latest_ter_url(html)
    assert str(d) == "2026-10-01" and url.endswith("TER_01.10.2026.xlsx")
    assert cf.latest_ter_url("nothing here") == (None, None)


def test_latest_riskometer_file_ignores_percent_20_in_urls():
    # regression: '%2020' + '26' used to be read as the year 2020 and the 2023 workbook won
    base = "https://assets-netstorage.growwmf.in/compliance_docs/Downloads/Riskometer/"
    html = " ".join(f'"{base}{n}"' for n in ["Riskometer-%202023-%202024.xlsx", "Riskometer%20-%202026%20-%202027.xlsx",
                                              "Riskometer%20-%202025-%202026.xlsx", "Riskometer%20-%202024-%202025.xlsx"])
    assert cf.latest_riskometer_url(html).endswith("Riskometer%20-%202026%20-%202027.xlsx")


@pytest.fixture
def network(monkeypatch, tmp_path):
    """Fake the network: links OK, live files = the committed copies (or tweaked)."""
    state = {"links_rc": 0, "table": "id   status via\nS01  200    HEAD application/pdf\n22/22 reachable", "ter": "data/raw/S10.xlsx"}
    monkeypatch.setattr(cf.check_links, "main", lambda: (print(state["table"]), state["links_rc"])[1])
    monkeypatch.setattr(cf, "page_text", lambda url: (
        '"https://assets-netstorage.growwmf.in/x/GRMF_WEBSITE%20TER_30.09.2026.xlsx"' if "expense" in url
        else '"https://assets-netstorage.growwmf.in/x/Riskometer/Riskometer%20-%202026%20-%202027.xlsx"'))
    monkeypatch.setattr(cf, "download", lambda url, suffix: state["ter"] if "TER" in url else "data/raw/S11.xlsx")
    monkeypatch.setenv("GITHUB_OUTPUT", str(tmp_path / "out.txt"))
    return state, tmp_path


def test_run_all_fresh_sets_no_problem_flag(network):
    state, tmp = network
    problems, report = cf.run(str(tmp / "r.md"))
    assert problems == [] and "All checks passed" in report
    assert "problems=false" in (tmp / "out.txt").read_text()


def test_run_with_a_broken_link_flags_it(network):
    state, tmp = network
    state["links_rc"], state["table"] = 1, "id   status via\nS10  404    GET  text/html\n21/22 reachable  - 1 PROBLEM(S)"
    problems, report = cf.run(str(tmp / "r.md"))
    assert any("Source link problem" in p and "S10" in p for p in problems)
    assert "problems=true" in (tmp / "out.txt").read_text()


def test_run_with_a_changed_ter_file_flags_it(network, tmp_path_factory):
    import openpyxl
    state, tmp = network
    wb = openpyxl.load_workbook("data/raw/S10.xlsx")
    ws = wb.worksheets[0]
    hdr = [c.value for c in ws[1]]
    col = hdr.index("Direct Plan - Total TER (%)") + 1
    for row in ws.iter_rows(min_row=2):
        if row[1].value == "Groww ELSS Tax Saver Fund":
            ws.cell(row=row[0].row, column=col).value = 1.99
    changed = tmp / "ter_changed.xlsx"
    wb.save(changed)
    state["ter"] = str(changed)
    problems, report = cf.run(str(tmp / "r.md"))
    assert any("ELSS Tax Saver Fund Direct plan is 1.99%" in p for p in problems) and "problem(s) found" in report
    assert "problems=true" in (tmp / "out.txt").read_text()

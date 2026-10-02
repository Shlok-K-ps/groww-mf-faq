"""Weekly data-freshness check (no secrets, no LLM calls).   python -m scripts.check_freshness [--report freshness_report.md]

1. Every URL in data/sources.csv still answers 200 (same check as scripts/check_links.py).
2. Fetches the LATEST expense-ratio (TER) file and riskometer file linked from growwmf.in, re-parses them, and compares each
   in-scope scheme's TER (Direct / Regular) and riskometer level with data/facts.json.

Any broken link or any TER/riskometer value that differs from facts.json is a PROBLEM: the report says what, and (when run in GitHub Actions)
sets `problems=true` so the workflow opens an issue. Exit code is 0 unless --fail-on-problems is given.
"""
import argparse
import contextlib
import datetime as dt
import html
import io
import json
import os
import re
import sys
import tempfile
from urllib.parse import unquote

import openpyxl
import requests

from rag.config import SCHEME_NAMES
from scripts import check_links

H = {"User-Agent": "Mozilla/5.0 (freshness check)"}
TER_PAGE = "https://www.growwmf.in/downloads/expense-ratio"
RISK_PAGE = "https://www.growwmf.in/downloads/riskometer"
TOL = 0.005


def page_text(url):
    t = requests.get(url, headers=H, timeout=60).text
    return html.unescape(t).replace("\\u002F", "/").replace("\\/", "/").replace("\\u0026", "&")


def latest_ter_url(text):
    """Newest 'GRMF_WEBSITE TER_DD.MM.YYYY.xlsx' linked from the page -> (url, date) or (None, None)."""
    found = []
    for m in re.finditer(r"https://assets-netstorage\.growwmf\.in/[^\"'\s\\<>]*?GRMF_WEBSITE%20TER_(\d{2})\.(\d{2})\.(\d{4})\.xlsx", text):
        d = dt.date(int(m.group(3)), int(m.group(2)), int(m.group(1)))
        found.append((d, m.group(0)))
    return (max(found)[1], max(found)[0]) if found else (None, None)


def latest_riskometer_url(text):
    """Newest yearly riskometer workbook ('.../Riskometer/Riskometer - 2026 - 2027.xlsx') -> url or None."""
    found = []
    for m in re.finditer(r"https://assets-netstorage\.growwmf\.in/[^\"'\s\\<>]*?/Riskometer/[^\"'\s\\<>]*?\.xlsx", text):
        years = tuple(int(y) for y in re.findall(r"(?<!\d)(20\d\d)(?!\d)", unquote(m.group(0).split("/Riskometer/")[1])))   # decode %20 first
        found.append((years, m.group(0)))
    return max(found)[1] if found else None


def _scheme(name):
    k = re.sub(r"[^a-z]", "", str(name or "").lower())
    return next((s for s in SCHEME_NAMES if re.sub(r"[^a-z]", "", s.lower()) == k), None)


def parse_ter(path):
    """{scheme: {"direct": float, "regular": float, "date": date}} for the four in-scope schemes."""
    ws = openpyxl.load_workbook(path, data_only=True).worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    hdr = [str(h).strip() if h else "" for h in rows[0]]
    ci = {"scheme": hdr.index("Scheme Name"), "date": hdr.index("TER Date"), "direct": hdr.index("Direct Plan - Total TER (%)"),
          "regular": hdr.index("Regular Plan - Total TER (%)")}
    out = {}
    for r in rows[1:]:
        s = _scheme(r[ci["scheme"]])
        if s:
            d = r[ci["date"]]
            out[s] = {"direct": round(float(r[ci["direct"]]), 2), "regular": round(float(r[ci["regular"]]), 2),
                      "date": d.date() if hasattr(d, "date") else None}
    return out


def parse_riskometer(path):
    """{scheme: (month_label, level)} using the latest month that has a value."""
    ws = openpyxl.load_workbook(path, data_only=True).worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    months = [(c, v) for c, v in enumerate(rows[1]) if hasattr(v, "year")]
    out = {}
    for r in rows[2:]:
        s = _scheme(r[2])
        if not s:
            continue
        vals = [(d, str(r[c]).strip()) for c, d in months if r[c] not in (None, "")]
        if vals:
            d, level = max(vals)
            out[s] = (d.strftime("%b %Y"), level)
    return out


def load_facts(path=os.path.join("data", "facts.json")):
    return json.load(open(path, encoding="utf-8"))


def compare(facts, ter, risk):
    """List of human-readable problems where a live value differs from facts.json (empty list = fresh)."""
    problems = []
    for s in SCHEME_NAMES:
        for plan in (("direct", "regular") if ter is not None else ()):
            f = next((x for x in facts if x["scheme"] == s and x["field"] == "expense_ratio" and x["plan"] == plan), None)
            live = (ter.get(s) or {}).get(plan)
            if f is None or live is None:
                problems.append(f"TER: {s} ({plan}) is missing from {'facts.json' if f is None else 'the live file'}")
            elif abs(float(f["value"]) - live) > TOL:
                problems.append(f"TER changed: {s} {plan.title()} plan is {live}% in the live file (as of {ter[s]['date']}) but facts.json has "
                                f"{f['value']}% (as of {f['effective_date']})")
        if risk is None:
            continue
        f = next((x for x in facts if x["scheme"] == s and x["field"] == "riskometer"), None)
        live = risk.get(s)
        if f is None or live is None:
            problems.append(f"Riskometer: {s} is missing from {'facts.json' if f is None else 'the live file'}")
        else:
            fact_month = f["conditions"].replace("Riskometer level for ", "")
            if live[1].lower() != f["value"].lower():
                problems.append(f"Riskometer changed: {s} is '{live[1]}' for {live[0]} in the live file but facts.json has "
                                f"'{f['value']}' for {fact_month}")
    return problems


def download(url, suffix):
    r = requests.get(url, headers=H, timeout=120)
    r.raise_for_status()
    fd, path = tempfile.mkstemp(suffix=suffix)
    with os.fdopen(fd, "wb") as f:
        f.write(r.content)
    return path


def run(report_path=None, facts=None):
    facts = facts if facts is not None else load_facts()
    problems, notes = [], []

    # 1. links
    buf = io.StringIO()
    with contextlib.redirect_stdout(buf):
        link_rc = check_links.main()
    table = buf.getvalue()
    bad_lines = [l for l in table.splitlines() if re.match(r"S\d+\s", l) and not re.match(r"S\d+\s+200\s", l)]
    if link_rc != 0 or bad_lines:
        problems += [f"Source link problem: {l.strip()}" for l in bad_lines] or ["Source link check failed (see log)"]
    notes.append(f"Link check: {table.strip().splitlines()[-1]}")

    # 2. TER + riskometer
    ter, risk = None, None            # None = file could not be fetched (already reported); {} = fetched but nothing parsed
    try:
        url, d = latest_ter_url(page_text(TER_PAGE))
        if not url:
            problems.append(f"Could not find a TER file link on {TER_PAGE} (page layout may have changed)")
        else:
            ter = parse_ter(download(url, ".xlsx"))
            notes.append(f"Latest TER file on growwmf.in: dated {d.isoformat()} ({os.path.basename(url.replace('%20', ' '))})")
    except Exception as e:
        problems.append(f"TER check failed: {type(e).__name__}: {e}")
    try:
        url = latest_riskometer_url(page_text(RISK_PAGE))
        if not url:
            problems.append(f"Could not find a riskometer file link on {RISK_PAGE} (page layout may have changed)")
        else:
            risk = parse_riskometer(download(url, ".xlsx"))
            latest = max((v[0] for v in risk.values()), key=lambda m: dt.datetime.strptime(m, "%b %Y"), default="none parsed")
            notes.append(f"Latest riskometer month on growwmf.in: {latest} ({unquote(os.path.basename(url))})")
    except Exception as e:
        problems.append(f"Riskometer check failed: {type(e).__name__}: {e}")
    problems += compare(facts, ter, risk)

    fact_dates = sorted({x["effective_date"] for x in facts if x["field"] in ("expense_ratio", "riskometer")})
    notes.append(f"facts.json TER/riskometer as-of dates: {', '.join(fact_dates)}")
    lines = [f"# Data freshness check - {dt.date.today().isoformat()}", "",
             ("**" + str(len(problems)) + " problem(s) found.** `data/facts.json` is out of date or a source is unreachable; "
              "re-run `python ingest.py` (parse, embed, facts, patch) and review.") if problems else "**All checks passed.** Links are reachable and "
             "the live TER and riskometer values match `data/facts.json`.", ""]
    if problems:
        lines += ["## Problems", ""] + [f"- {p}" for p in problems] + [""]
    lines += ["## Details", ""] + [f"- {n}" for n in notes] + ["", "<details><summary>Link check output</summary>", "", "```", table.strip(), "```", "", "</details>", ""]
    report = "\n".join(lines)
    if report_path:
        open(report_path, "w", encoding="utf-8").write(report)
    gh_out = os.environ.get("GITHUB_OUTPUT")
    if gh_out:
        with open(gh_out, "a", encoding="utf-8") as f:
            f.write(f"problems={'true' if problems else 'false'}\n")
    summary = os.environ.get("GITHUB_STEP_SUMMARY")
    if summary:
        open(summary, "a", encoding="utf-8").write(report + "\n")
    return problems, report


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--report", default="freshness_report.md")
    ap.add_argument("--fail-on-problems", action="store_true")
    a = ap.parse_args()
    sys.stdout.reconfigure(encoding="utf-8")
    problems, report = run(a.report)
    print(report)
    return 1 if (problems and a.fail_on_problems) else 0


if __name__ == "__main__":
    sys.exit(main())
